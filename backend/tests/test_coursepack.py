"""Course upload: local text extraction, verbatim classify, pack of practices. No network."""
import io
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
os.environ["GATEWAY_API_KEY"] = ""
os.environ["LLM_API_KEY"] = ""

from app import config, db  # noqa: E402
from app.coursepack import normalize, save_course  # noqa: E402
from app.materials import ExtractError, extract  # noqa: E402

SOURCE = """
--- page 4 ---
Introduce yourself!
My name is ________. I come from ________.
--- page 8 ---
What's the main business of your company?
Warm up
--- page 10 ---
specialize
to be an expert in a particular part of a subject
What does your company specialize in?
--- page 17 ---
What kind of company is it?
We are a construction company.
--- page 20 ---
You meet a friend whom you haven't seen in a while. You both have your own businesses.
--- page 24 ---
Matt: Hi, Alex! I heard you started your own business.
Alex: We design software applications for smartphones.
--- page 25 ---
What kind of business did Alex start?
"""


def fresh_db():
    config.DB_PATH = Path(tempfile.mkdtemp()) / "t.db"
    db.init()


def test_extract_pdf_docx_pptx_locally():
    from pypdf import PdfWriter
    from docx import Document
    from pptx import Presentation

    doc = Document()
    doc.add_paragraph("My name is ________. I come from ________. What does your company specialize in?")
    buf = io.BytesIO()
    doc.save(buf)
    text = extract("notes.docx", buf.getvalue())
    assert "specialize" in text and "--- page" not in text

    pres = Presentation()
    slide = pres.slides.add_slide(pres.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, 4000000, 800000)
    box.text_frame.text = "What's the main business of your company?"
    buf = io.BytesIO()
    pres.save(buf)
    text = extract("deck.pptx", buf.getvalue())
    assert "--- slide 1 ---" in text and "main business" in text

    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    buf = io.BytesIO()
    writer.write(buf)
    try:
        extract("blank.pdf", buf.getvalue())
        raise AssertionError("blank pdf should fail")
    except ExtractError as e:
        assert "没有提取到足够的文字" in str(e)
    try:
        extract("old.ppt", b"not a ppt")
        raise AssertionError("ppt should fail")
    except ExtractError as e:
        assert "pptx" in str(e)


def test_normalize_keeps_only_verbatim_lines():
    data = {
        "title": "Introduce yourself!",
        "topic": "描述业务",
        "summary": "介绍公司业务",
        "questions": [
            {"text": "What's the main business of your company?", "tip": "说说主营业务"},
            {"text": "What is your favourite planet?", "tip": "这句是编的"},
        ],
        "vocabulary": [{
            "word": "specialize", "pos": "", "meaning": "to be an expert in a particular part of a subject",
            "example": "What does your company specialize in?",
        }],
        "patterns": [{"pattern": "What kind of company is it?", "usage": "问公司类型", "examples": [
            {"text": "We are a construction company.", "translation": ""},
        ]}],
        "introduce": {"template": "My name is ________. I come from ________."},
        "chat": {"scene": "多年没见的朋友，各自做生意", "goal": "聊聊彼此的工作", "aiRole": "老朋友", "opening": ""},
        "dialogue": {
            "roles": ["Matt", "Alex"],
            "lines": [
                {"role": "Matt", "text": "Hi, Alex! I heard you started your own business."},
                {"role": "Alex", "text": "We design software applications for smartphones."},
                {"role": "Matt", "text": "I invented a new product line yesterday."},
            ],
        },
        "listening": {"questions": [{"text": "What kind of business did Alex start?", "answerHint": ""}]},
    }
    pack = normalize(data, SOURCE)
    assert [q["text"] for q in pack["questions"]] == ["What's the main business of your company?"]
    assert pack["vocabulary"][0]["word"] == "specialize"
    assert pack["template"].startswith("My name is")
    assert len(pack["dialogue"]["lines"]) == 2
    assert pack["listeningQuestions"][0]["text"].startswith("What kind of business")


def test_save_course_builds_practices_and_notes():
    fresh_db()
    data = {
        "title": "Introduce yourself!",
        "topic": "描述业务",
        "summary": "介绍公司业务",
        "questions": [
            {"text": "What's the main business of your company?", "tip": "说说主营业务"},
            {"text": "What does your company specialize in?", "tip": ""},
        ],
        "vocabulary": [{
            "word": "specialize", "pos": "", "meaning": "to be an expert in a particular part of a subject",
            "example": "What does your company specialize in?",
        }],
        "patterns": [{"pattern": "What kind of company is it?", "usage": "问公司类型", "examples": []}],
        "introduce": {"template": "My name is ________. I come from ________."},
        "chat": {"scene": "多年没见的朋友，各自做生意", "goal": "聊聊工作", "aiRole": "老朋友", "opening": ""},
        "dialogue": {"roles": ["Matt", "Alex"], "lines": [
            {"role": "Matt", "text": "Hi, Alex! I heard you started your own business."},
            {"role": "Alex", "text": "We design software applications for smartphones."},
        ]},
        "listening": {"lines": [
            {"role": "Matt", "text": "Hi, Alex! I heard you started your own business."},
            {"role": "Alex", "text": "We design software applications for smartphones."},
        ], "questions": [{"text": "What kind of business did Alex start?", "answerHint": ""}]},
    }
    saved = save_course(1, "en", "intermediate", "business.pdf", SOURCE, data)
    course = db.get_course(saved["id"])
    kinds = [it["kind"] for it in course["items"]]
    assert kinds == ["questions", "notes", "dictation", "chat", "dialogue", "listening", "introduce"]
    notes_item = course["items"][1]
    assert notes_item["sessionId"] == course["items"][0]["sessionId"]
    host = db.get_session(notes_item["sessionId"])["content"]
    assert host.get("notesLocked") is True
    notes = db.get_notes(notes_item["sessionId"])["notes"]
    assert notes["vocabulary"][0]["word"] == "specialize"
    assert notes["patterns"][0]["pattern"] == "What kind of company is it?"
    dictation = db.get_session(course["items"][2]["sessionId"])["content"]
    assert dictation["sentences"][0]["blanks"] == ["specialize"]
    dialogue = db.get_session(course["items"][4]["sessionId"])["content"]
    assert {r["name"] for r in dialogue["roles"]} == {"Matt", "Alex"}
    chat = db.get_session(course["items"][3]["sessionId"])["content"]
    assert chat["opening"] == "Hi, Alex! I heard you started your own business."
    assert db.get_course("nope") is None


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn()
            print("ok", name)
