"""Pull plain text out of an uploaded course file. No model call."""
import io
from pathlib import Path

MAX_BYTES = 20 * 1024 * 1024
MAX_CHARS = 40_000

class ExtractError(ValueError):
    pass


def extract(filename: str, data: bytes) -> str:
    if not data:
        raise ExtractError("文件是空的")
    if len(data) > MAX_BYTES:
        raise ExtractError("文件超过 20MB，请拆开或改成粘贴文字")
    ext = Path(filename or "").suffix.lower()
    if ext == ".pdf":
        text = _pdf(data)
    elif ext == ".pptx":
        text = _pptx(data)
    elif ext == ".docx":
        text = _docx(data)
    elif ext == ".ppt":
        raise ExtractError("旧版 .ppt 读不了，请在 PowerPoint 里另存为 .pptx")
    elif ext == ".doc":
        raise ExtractError("旧版 .doc 读不了，请在 Word 里另存为 .docx")
    elif ext in {".md", ".markdown", ".txt", ".zip"}:
        raise ExtractError("这份文件请用管理员账号在首页「上传资料」里导入；整理课程包只收 PDF、PPTX、DOCX，或直接粘贴文字")
    else:
        raise ExtractError("只支持 PDF、PPTX、DOCX，或直接粘贴文字")
    text = text.strip()
    if len(text) < 40:
        raise ExtractError("没有提取到足够的文字。扫描版 PDF 请换一份可选中文字的文件，或直接粘贴。")
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS]
    return text


def _pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
    except Exception as e:
        raise ExtractError(f"这份 PDF 打不开：{e}") from e
    parts = []
    for i, page in enumerate(reader.pages, 1):
        body = (page.extract_text() or "").strip()
        if body:
            parts.append(f"--- page {i} ---\n{body}")
    return "\n\n".join(parts)


def _pptx(data: bytes) -> str:
    try:
        from pptx import Presentation
        pres = Presentation(io.BytesIO(data))
    except Exception as e:
        raise ExtractError(f"这份 PPT 打不开：{e}") from e
    parts = []
    for i, slide in enumerate(pres.slides, 1):
        lines = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                chunk = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
                if chunk.strip():
                    lines.append(chunk.strip())
        if lines:
            parts.append(f"--- slide {i} ---\n" + "\n".join(lines))
    return "\n\n".join(parts)


def _docx(data: bytes) -> str:
    try:
        from docx import Document
        doc = Document(io.BytesIO(data))
    except Exception as e:
        raise ExtractError(f"这份 Word 打不开：{e}") from e
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)
